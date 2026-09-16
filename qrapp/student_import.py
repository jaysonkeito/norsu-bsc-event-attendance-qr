import csv
import os
import re

import pdfplumber

HEADER_ALIASES = {
    "student_id": [
        "student id",
        "student no id",
        "student no",
        "id number",
        "id",
        "student_id",
        "student number",
        "id no",
    ],
    "name": ["name", "full name", "student name"],
    "sex": ["sex", "gender"],
    "college": ["college"],
    "program": ["program", "course"],
    "year": ["year", "year level", "yr", "yearlevel"],
    "major": ["major", "section", "sec"],
}

SKIP_ROW_MARKERS = {"student id", "id number", "generated", "id no", "student no"}


def safe_strip(value, default=""):
    if value is None:
        return default
    value_str = str(value).strip()
    return value_str if value_str else default


def normalize_header(value):
    normalized = safe_strip(value).lower().replace("_", " ")
    normalized = re.sub(r"[./]+", " ", normalized)
    return re.sub(r"\s+", " ", normalized).strip()


def is_header_row(row):
    if not row:
        return False
    first = normalize_header(row[0])
    if first in SKIP_ROW_MARKERS:
        return True
    normalized = [normalize_header(cell) for cell in row if safe_strip(cell)]
    for aliases in HEADER_ALIASES.values():
        if any(alias in normalized for alias in aliases):
            return True
    return False


def build_column_map(header_row):
    column_map = {}
    for index, cell in enumerate(header_row):
        header = normalize_header(cell)
        if not header:
            continue
        for field, aliases in HEADER_ALIASES.items():
            if header in aliases and field not in column_map:
                column_map[field] = index
                break
    return column_map


def parse_student_row(row, column_map=None):
    if not row:
        return None

    first_cell = normalize_header(row[0])
    if first_cell in SKIP_ROW_MARKERS:
        return None

    if column_map:
        def get_field(field, default=""):
            index = column_map.get(field)
            if index is None or index >= len(row):
                return default
            return safe_strip(row[index], default)

        student_id = get_field("student_id", "0")
        name = get_field("name")
        sex = get_field("sex")
        college = get_field("college", "CAS")
        program = get_field("program")
        year = get_field("year", "0")
        major = get_field("major", "NA")
    else:
        # The first column is a spreadsheet row number, not the student ID.
        student_id = safe_strip(row[1], "0") if len(row) > 1 else "0"
        name = safe_strip(row[2]) if len(row) > 2 else ""
        sex = safe_strip(row[3]) if len(row) > 3 else ""
        college = "CAS"
        program = safe_strip(row[4]) if len(row) > 4 else ""
        year = safe_strip(row[5], "0") if len(row) > 5 else "0"
        major = safe_strip(row[6], "NA") if len(row) > 6 else "NA"

    if student_id in ("0", "NA", "") or name in ("NA", ""):
        return None

    try:
        year_int = int(float(str(year).replace(",", "")))
    except (ValueError, TypeError):
        year_int = 0

    return {
        "student_id": student_id,
        "name": name,
        "sex": sex,
        "college": college,
        "program": program,
        "year": year_int,
        "major": major,
    }


def import_students_from_rows(rows, column_map=None):
    from .college_program import resolve_college_and_program
    from .models import Student

    created = 0
    skipped = 0

    for row in rows:
        data = parse_student_row(row, column_map)
        if not data:
            skipped += 1
            continue

        college, program = resolve_college_and_program(data["college"], data["program"])

        _, was_created = Student.objects.get_or_create(
            student_id=data["student_id"],
            defaults={
                "name": data["name"],
                "sex": data["sex"],
                "college": college,
                "program": program,
                "year": data["year"],
                "major": data["major"],
            },
        )
        if was_created:
            created += 1
        else:
            skipped += 1

    return created, skipped


def parse_pdf_text_row(line):
    tokens = safe_strip(line).split()
    if len(tokens) < 5:
        return None

    sex_index = next(
        (index for index, token in enumerate(tokens) if token.lower() in {"m", "f", "male", "female"}),
        None,
    )
    if sex_index is None or sex_index < 2:
        return None

    year_index = next(
        (
            index
            for index in range(sex_index + 1, len(tokens))
            if re.fullmatch(r"\d{1,2}", tokens[index])
        ),
        None,
    )
    if year_index is None or year_index <= sex_index + 1:
        return None

    # A row number may precede the student ID in the requested PDF format.
    id_index = 1 if tokens[0].isdigit() and sex_index >= 3 else 0
    if id_index >= sex_index - 1:
        return None

    return [
        tokens[id_index],
        " ".join(tokens[id_index + 1:sex_index]),
        tokens[sex_index],
        " ".join(tokens[sex_index + 1:year_index]),
        tokens[year_index],
        " ".join(tokens[year_index + 1:]),
    ]


def iter_pdf_rows(file_path):
    rows = []
    column_map = None
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            table = page.extract_table()
            if not table:
                continue
            if is_header_row(table[0]):
                if column_map is None:
                    column_map = build_column_map(table[0])
                rows.extend(table[1:])
            else:
                rows.extend(table)

        if rows:
            return rows, column_map

        text_column_map = {
            "student_id": 0,
            "name": 1,
            "sex": 2,
            "program": 3,
            "year": 4,
            "major": 5,
        }
        for page in pdf.pages:
            text = page.extract_text() or ""
            for line in text.splitlines():
                if is_header_row(line.split()) or normalize_header(line).startswith("generated"):
                    continue
                parsed_row = parse_pdf_text_row(line)
                if parsed_row:
                    rows.append(parsed_row)
        if rows:
            return rows, text_column_map
    return rows, column_map


def iter_csv_rows(file_path):
    rows = []
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            with open(file_path, newline="", encoding=encoding) as handle:
                reader = csv.reader(handle)
                for row in reader:
                    if any(safe_strip(cell) for cell in row):
                        rows.append(row)
            break
        except UnicodeDecodeError:
            rows = []
            continue

    if not rows:
        raise ValueError("Could not read CSV file. Please save it as UTF-8.")

    column_map = None
    data_rows = rows
    if is_header_row(rows[0]):
        column_map = build_column_map(rows[0])
        data_rows = rows[1:]

    return data_rows, column_map


def iter_xlsx_rows(file_path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("Excel support requires openpyxl. Run: pip install openpyxl") from exc

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = []
    for row in sheet.iter_rows(values_only=True):
        row_values = [cell if cell is not None else "" for cell in row]
        if any(safe_strip(cell) for cell in row_values):
            rows.append(row_values)
    workbook.close()

    if not rows:
        return [], None

    column_map = None
    data_rows = rows
    if is_header_row(rows[0]):
        column_map = build_column_map(rows[0])
        data_rows = rows[1:]

    return data_rows, column_map


def import_students_from_file(file_path):
    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        rows, column_map = iter_pdf_rows(file_path)
        label = "PDF"
    elif extension == ".csv":
        rows, column_map = iter_csv_rows(file_path)
        label = "CSV"
    elif extension in (".xlsx", ".xlsm"):
        rows, column_map = iter_xlsx_rows(file_path)
        label = "Excel"
    else:
        raise ValueError("Unsupported file type. Use PDF, CSV, or Excel (.xlsx).")

    if not rows:
        raise ValueError(f"No student rows found in the {label} file.")

    created, skipped = import_students_from_rows(rows, column_map)
    return {
        "created": created,
        "skipped": skipped,
        "label": label,
        "message": f"Added {created} students from {label}! ({skipped} rows skipped or already exist)",
    }
